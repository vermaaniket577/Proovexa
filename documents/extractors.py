"""
Text extraction engine for PDF, DOCX, TXT, and PPTX files.
Preserves page numbers and paragraph boundaries.
"""
import os
import re
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class ExtractionResult:
    """Result of text extraction from a document."""

    def __init__(self):
        self.pages = []  # List of dicts: {page_number, paragraphs: [{text, para_number}]}
        self.full_text = ""
        self.word_count = 0
        self.page_count = 0
        self.errors = []

    def add_page(self, page_number, paragraphs):
        """Add a page with its paragraphs."""
        page_data = {
            'page_number': page_number,
            'paragraphs': []
        }
        for i, text in enumerate(paragraphs, 1):
            cleaned = text.strip()
            if cleaned:
                page_data['paragraphs'].append({
                    'text': cleaned,
                    'para_number': i,
                })
        if page_data['paragraphs']:
            self.pages.append(page_data)

    def finalize(self):
        """Compute final statistics."""
        self.page_count = len(self.pages)
        parts = []
        for page in self.pages:
            for para in page['paragraphs']:
                parts.append(para['text'])
        self.full_text = '\n\n'.join(parts)
        self.word_count = len(self.full_text.split())
        return self


class TextExtractor:
    """Base text extractor."""

    def extract(self, file_path):
        raise NotImplementedError


class PDFExtractor(TextExtractor):
    """Extract text from PDF files using PyMuPDF."""

    def extract(self, file_path):
        import fitz  # PyMuPDF

        result = ExtractionResult()
        try:
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text("text")
                if text:
                    # Split into paragraphs by double newlines or significant breaks
                    paragraphs = self._split_paragraphs(text)
                    result.add_page(page_num + 1, paragraphs)
            doc.close()
        except Exception as e:
            logger.error(f"PDF extraction error: {e}")
            result.errors.append(f"PDF extraction error: {str(e)}")

        return result.finalize()

    def _split_paragraphs(self, text):
        """Split text into paragraphs."""
        # Split on double newlines or multiple whitespace lines
        paragraphs = re.split(r'\n\s*\n', text)
        # Clean each paragraph
        cleaned = []
        for p in paragraphs:
            p = ' '.join(p.split())  # Normalize whitespace
            if len(p.strip()) > 0:
                cleaned.append(p)
        return cleaned


class DOCXExtractor(TextExtractor):
    """Extract text from DOCX files using python-docx."""

    def extract(self, file_path):
        from docx import Document as DocxDocument

        result = ExtractionResult()
        try:
            doc = DocxDocument(file_path)
            current_paragraphs = []
            current_page = 1

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                # Check for page breaks in paragraph's runs
                has_page_break = False
                for run in para.runs:
                    if run._element.xml.find('w:br') != -1 and 'type="page"' in run._element.xml:
                        has_page_break = True
                        break

                if has_page_break and current_paragraphs:
                    result.add_page(current_page, current_paragraphs)
                    current_paragraphs = []
                    current_page += 1

                current_paragraphs.append(text)

            # Add last page
            if current_paragraphs:
                result.add_page(current_page, current_paragraphs)

        except Exception as e:
            logger.error(f"DOCX extraction error: {e}")
            result.errors.append(f"DOCX extraction error: {str(e)}")

        return result.finalize()


class TXTExtractor(TextExtractor):
    """Extract text from plain text files."""

    def extract(self, file_path):
        result = ExtractionResult()
        try:
            encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']
            text = None
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        text = f.read()
                    break
                except (UnicodeDecodeError, UnicodeError):
                    continue

            if text is None:
                result.errors.append("Could not decode the text file with supported encodings.")
                return result.finalize()

            # Split into paragraphs
            paragraphs = re.split(r'\n\s*\n', text)
            cleaned = [p.strip() for p in paragraphs if p.strip()]

            # Treat every ~50 paragraphs as a "page" for long documents
            page_size = 50
            for i in range(0, len(cleaned), page_size):
                page_paragraphs = cleaned[i:i + page_size]
                result.add_page((i // page_size) + 1, page_paragraphs)

        except Exception as e:
            logger.error(f"TXT extraction error: {e}")
            result.errors.append(f"TXT extraction error: {str(e)}")

        return result.finalize()


class PPTXExtractor(TextExtractor):
    """Extract text from PPTX files using python-pptx."""

    def extract(self, file_path):
        from pptx import Presentation

        result = ExtractionResult()
        try:
            prs = Presentation(file_path)
            for slide_num, slide in enumerate(prs.slides, 1):
                paragraphs = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        paragraphs.append(shape.text.strip())
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            row_text = ' | '.join(
                                cell.text.strip() for cell in row.cells if cell.text.strip()
                            )
                            if row_text:
                                paragraphs.append(row_text)
                if paragraphs:
                    result.add_page(slide_num, paragraphs)

        except Exception as e:
            logger.error(f"PPTX extraction error: {e}")
            result.errors.append(f"PPTX extraction error: {str(e)}")

        return result.finalize()


def get_extractor(file_type):
    """Factory function to get the appropriate extractor."""
    extractors = {
        'pdf': PDFExtractor,
        'docx': DOCXExtractor,
        'txt': TXTExtractor,
        'pptx': PPTXExtractor,
    }
    extractor_class = extractors.get(file_type.lower())
    if extractor_class is None:
        raise ValueError(f"Unsupported file type: {file_type}")
    return extractor_class()


def extract_text(file_path, file_type):
    """
    Extract text from a document file.
    
    Args:
        file_path: Path to the file
        file_type: File type (pdf, docx, txt, pptx)
    
    Returns:
        ExtractionResult with pages, paragraphs, and statistics
    """
    extractor = get_extractor(file_type)
    return extractor.extract(file_path)
