"""
DOCX report generator using python-docx.
Creates professional academic similarity reports.
"""
import io
from datetime import datetime
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


# Source colors for highlighting
SOURCE_COLORS_RGB = [
    RGBColor(231, 76, 60),   # Red
    RGBColor(52, 152, 219),  # Blue
    RGBColor(46, 204, 113),  # Green
    RGBColor(243, 156, 18),  # Orange
    RGBColor(155, 89, 182),  # Purple
    RGBColor(26, 188, 156),  # Teal
    RGBColor(230, 126, 34),  # Dark Orange
    RGBColor(52, 73, 94),    # Dark Gray
    RGBColor(192, 57, 43),   # Dark Red
    RGBColor(41, 128, 185),  # Dark Blue
]


def generate_docx_report(report):
    """
    Generate a professional DOCX similarity report.
    
    Args:
        report: SimilarityReport model instance
    
    Returns:
        BytesIO buffer containing the DOCX
    """
    doc = DocxDocument()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # ---- Title ----
    title = doc.add_heading('Similarity Analysis Report', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(26, 26, 46)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        f"Generated on {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    )
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(102, 102, 102)

    doc.add_paragraph('_' * 80)

    # ---- Document Information ----
    doc.add_heading('Document Information', level=2)

    info_table = doc.add_table(rows=8, cols=2)
    info_table.style = 'Light Grid Accent 1'

    info_data = [
        ('Report ID', str(report.id)[:8].upper()),
        ('Document ID', str(report.document.id)[:8].upper()),
        ('Document Title', report.document.title),
        ('Author', report.document.uploaded_by.get_full_name() or report.document.uploaded_by.username),
        ('File Type', report.document.get_file_type_display()),
        ('Upload Date', report.document.uploaded_at.strftime('%Y-%m-%d %H:%M')),
        ('Total Words', f"{report.total_words:,}"),
        ('Pages', str(report.document.page_count)),
    ]

    for i, (label, value) in enumerate(info_data):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = value
        # Bold the labels
        for paragraph in info_table.rows[i].cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    doc.add_paragraph()

    # ---- Similarity Score ----
    doc.add_heading('Similarity Results', level=2)

    score_para = doc.add_paragraph()
    score_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    score_run = score_para.add_run(f"{report.overall_similarity:.1f}%")
    score_run.font.size = Pt(36)
    score_run.bold = True
    score_run.font.color.rgb = _get_score_color_rgb(report.overall_similarity)

    label_para = doc.add_paragraph()
    label_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_run = label_para.add_run('Overall Similarity Score')
    label_run.font.size = Pt(10)
    label_run.font.color.rgb = RGBColor(102, 102, 102)

    # Results table
    results_table = doc.add_table(rows=8, cols=2)
    results_table.style = 'Light Grid Accent 1'

    results_data = [
        ('Matched Words', f"{report.matched_words:,} of {report.total_words:,}"),
        ('Matching Sources', str(report.source_count)),
        ('Exact Matches', str(report.exact_match_count)),
        ('Phrase Matches', str(report.phrase_match_count)),
        ('Fuzzy Matches', str(report.fuzzy_match_count)),
        ('Processing Time', f"{report.processing_time_seconds:.1f} seconds"),
        ('Engine Version', report.engine_version),
        ('Analysis Date', report.created_at.strftime('%Y-%m-%d %H:%M:%S')),
    ]

    for i, (label, value) in enumerate(results_data):
        results_table.rows[i].cells[0].text = label
        results_table.rows[i].cells[1].text = value
        for paragraph in results_table.rows[i].cells[0].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    doc.add_paragraph()

    # ---- Source Breakdown ----
    matches = report.matched_passages.filter(is_excluded=False).select_related('source')
    source_groups = {}
    for match in matches:
        src_name = match.source.title if match.source else 'Unknown Source'
        if src_name not in source_groups:
            source_groups[src_name] = {
                'matches': 0, 'words': 0, 'types': set(), 'index': match.source_index
            }
        source_groups[src_name]['matches'] += 1
        source_groups[src_name]['words'] += match.word_count
        source_groups[src_name]['types'].add(match.get_match_type_display())

    if source_groups:
        doc.add_heading('Source-by-Source Breakdown', level=2)
        src_table = doc.add_table(
            rows=len(source_groups) + 1, cols=5
        )
        src_table.style = 'Light Grid Accent 1'

        headers = ['#', 'Source', 'Matches', 'Words', 'Match Types']
        for i, header in enumerate(headers):
            src_table.rows[0].cells[i].text = header
            for para in src_table.rows[0].cells[i].paragraphs:
                for run in para.runs:
                    run.bold = True

        for row_idx, (name, data) in enumerate(sorted(
            source_groups.items(), key=lambda x: x[1]['words'], reverse=True
        ), 1):
            src_table.rows[row_idx].cells[0].text = str(row_idx)
            src_table.rows[row_idx].cells[1].text = name[:50]
            src_table.rows[row_idx].cells[2].text = str(data['matches'])
            src_table.rows[row_idx].cells[3].text = str(data['words'])
            src_table.rows[row_idx].cells[4].text = ', '.join(data['types'])

    doc.add_paragraph()

    # ---- Matched Passages ----
    doc.add_page_break()
    doc.add_heading('Matched Passages', level=2)

    non_excluded = matches.order_by('submitted_page', 'submitted_char_start')
    if non_excluded.exists():
        for i, match in enumerate(non_excluded[:100], 1):
            color = SOURCE_COLORS_RGB[(match.source_index - 1) % len(SOURCE_COLORS_RGB)]

            # Match header
            header_para = doc.add_paragraph()
            run = header_para.add_run(f'Match {i}')
            run.bold = True
            run.font.color.rgb = color
            run = header_para.add_run(
                f' | {match.get_match_type_display()} | '
                f'Page {match.submitted_page}, Para {match.submitted_paragraph} | '
                f'Similarity: {match.similarity_score:.0%}'
            )
            run.font.size = Pt(9)

            # Source info
            if match.source:
                src_para = doc.add_paragraph()
                run = src_para.add_run(f'Source: {match.source.title}')
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(102, 102, 102)

            # Submitted text
            text_para = doc.add_paragraph()
            text_para.paragraph_format.left_indent = Inches(0.3)
            submitted_preview = match.submitted_text[:500]
            if len(match.submitted_text) > 500:
                submitted_preview += '...'
            run = text_para.add_run(f'"{submitted_preview}"')
            run.italic = True
            run.font.color.rgb = color
            run.font.size = Pt(10)

            # Source text
            if match.source_text:
                src_text_para = doc.add_paragraph()
                src_text_para.paragraph_format.left_indent = Inches(0.3)
                source_preview = match.source_text[:500]
                if len(match.source_text) > 500:
                    source_preview += '...'
                run = src_text_para.add_run(f'Source text: "{source_preview}"')
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(136, 136, 136)

            doc.add_paragraph()
    else:
        doc.add_paragraph('No matching passages found.')

    # ---- Exclusion Settings ----
    doc.add_heading('Exclusion Settings', level=2)
    excl = report.exclusion_settings or {}
    excl_items = []
    if excl.get('exclude_bibliography'):
        excl_items.append('Bibliography/references excluded')
    if excl.get('exclude_quotations'):
        excl_items.append('Direct quotations excluded')
    if excl.get('exclude_citations'):
        excl_items.append('Formatted citations excluded')
    excl_items.append(f"Minimum match length: {excl.get('min_match_words', 14)} words")
    if excl.get('exclude_common_phrases'):
        excl_items.append('Common phrases excluded')
    if excl.get('exclude_self_match'):
        excl_items.append('Self-matching excluded')
    for item in excl_items:
        doc.add_paragraph(item, style='List Bullet')

    # ---- Disclaimer ----
    doc.add_paragraph()
    doc.add_paragraph('_' * 80)
    disclaimer_heading = doc.add_paragraph()
    run = disclaimer_heading.add_run('Disclaimer')
    run.bold = True
    run.font.color.rgb = RGBColor(192, 57, 43)

    disclaimer_para = doc.add_paragraph()
    run = disclaimer_para.add_run(report.limitations)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(102, 102, 102)

    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _get_score_color_rgb(score):
    """Get RGB color for similarity score."""
    if score < 15:
        return RGBColor(39, 174, 96)   # Green
    elif score < 25:
        return RGBColor(41, 128, 185)  # Blue
    elif score < 50:
        return RGBColor(243, 156, 18)  # Orange
    else:
        return RGBColor(231, 76, 60)   # Red
